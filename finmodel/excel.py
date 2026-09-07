"""Write engine outputs to .xlsx.  The three-statement and DCF writers emit LIVE formulas in the
CFI layout (blue = hard-coded inputs, black = formulas), so the workbook stays editable in Excel /
LibreOffice; the projection writer emits values.  `recalc_with_libreoffice` re-computes a workbook
headlessly so tests can prove the formulas reproduce the engine."""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .three_statement import ThreeStatementResult, HistoricalYear, ForecastAssumptions
from .dcf import DCFInputs

BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="999999")
NUM = "#,##0;(#,##0)"
NUM2 = "#,##0.00;(#,##0.00)"
PCT = "0.0%"


def _title(ws, cell, text):
    ws[cell] = text
    ws[cell].font = Font(bold=True, size=14)


def _section(ws, row, text, first_col=1, last_col=14):
    ws.cell(row=row, column=first_col, value=text).font = BOLD
    for c in range(first_col, last_col + 1):
        ws.cell(row=row, column=c).fill = SECTION_FILL


# --------------------------------------------------------------------------- three statement
def write_three_statement(path: str | Path, hist: Sequence[HistoricalYear], fc: ForecastAssumptions,
                          ppe_opening0: Optional[float] = None, debt_opening0: Optional[float] = None,
                          result: Optional[ThreeStatementResult] = None) -> Path:
    """Formula-driven workbook with the CFI 'Three Statement Model' row layout."""
    from . import three_statement as ts
    if result is None:
        result = ts.run(hist, fc, ppe_opening0, debt_opening0)
    A = fc.expanded()
    n_h, n_f = len(hist), fc.years
    N = n_h + n_f
    if ppe_opening0 is None:
        ppe_opening0 = hist[0].ppe - hist[0].capex + hist[0].da
    if debt_opening0 is None:
        debt_opening0 = hist[0].debt - hist[0].debt_issued

    wb = Workbook()
    ws = wb.active
    ws.title = "Three Statement Model"
    ws.column_dimensions["A"].width = 44
    first = 4  # column D
    cols = [get_column_letter(first + i) for i in range(N)]

    def put(row, i, value, fmt=NUM, is_input=False):
        c = ws.cell(row=row, column=first + i, value=value)
        c.number_format = fmt
        if is_input:
            c.font = BLUE
        return c

    # headers
    ws["D1"], ws[f"{cols[n_h]}1"] = "Historical Results", "Forecast Period"
    ws["D1"].font = BOLD; ws[f"{cols[n_h]}1"].font = BOLD
    for i, y in enumerate(result.years):
        c = put(2, i, y if i == 0 else f"={cols[i-1]}2+1", "0", is_input=(i == 0))
        c.font = HEAD; c.fill = HEAD_FILL; c.alignment = Alignment(horizontal="center")
    ws["A2"] = "USD $000"; ws["A2"].font = HEAD; ws["A2"].fill = HEAD_FILL
    ws["A3"] = "Balance Sheet Check"
    for i, c in enumerate(cols):
        ws[f"{c}3"] = f'=IFERROR(IF(ABS({c}60)>1,"ERROR","OK"),"OK")'

    # assumptions block rows 8-21
    _section(ws, 5, "Assumptions")
    ws["A7"] = "Key Assumptions"; ws["A7"].font = BOLD
    asm_rows = [
        (8, "Revenue Growth (% Change)", "revenue_growth", PCT, lambda c, p: f"={c}26/{p}26-1"),
        (9, "Cost of Goods Sold (% of Revenue)", "cogs_pct", PCT, lambda c, p: f"={c}27/{c}26"),
        (10, "Salaries and Benefits (% of Revenue)", "salaries_pct", PCT, lambda c, p: f"={c}30/{c}26"),
        (11, "Rent and Overhead ($000's)", "rent", NUM, lambda c, p: f"={c}31"),
        (12, "Depreciation & Amortization (% of PP&E)", "da_pct_ppe", PCT, lambda c, p: f"={c}32/{c}47"),
        (13, "Interest (% of Debt)", "interest_pct_debt", PCT, lambda c, p: f"={c}33/{c}52"),
        (14, "Tax Rate (% of Earnings Before Tax)", "tax_rate", PCT, lambda c, p: f"={c}37/{c}35"),
        (15, "Accounts Receivable (Days)", "ar_days", NUM2, lambda c, p: f"={c}45/{c}26*{c}21"),
        (16, "Inventory (Days)", "inventory_days", NUM2, lambda c, p: f"={c}46/{c}27*{c}21"),
        (17, "Accounts Payable (Days)", "ap_days", NUM2, lambda c, p: f"={c}51/{c}27*{c}21"),
        (18, "Capital Expenditures ($000's)", "capex", NUM, lambda c, p: f"={c}72"),
        (19, "Debt Issuance (Repayment) ($000's)", "debt_issued", NUM, lambda c, p: f"={c}76"),
        (20, "Equity Issued (Repaid) ($000's)", "equity_issued", NUM, lambda c, p: f"={c}77"),
        (21, "Days in Period", "days_in_period", "0", None),
    ]
    for row, label, key, fmt, hist_formula in asm_rows:
        ws[f"A{row}"] = label
        for i, c in enumerate(cols):
            p = cols[i - 1] if i else None
            if i < n_h:
                if key == "days_in_period":
                    put(row, i, 365, fmt, is_input=True)
                elif row == 8 and i == 0:
                    continue
                else:
                    put(row, i, hist_formula(c, p), fmt)
            else:
                put(row, i, A[key][i - n_h], fmt, is_input=True)

    # income statement rows 26-38
    _section(ws, 24, "Income Statement")
    is_rows = {26: "Revenue", 27: "Cost of Goods Sold (COGS)", 28: "Gross Profit", 29: "Expenses",
               30: "Salaries and Benefits", 31: "Rent and Overhead", 32: "Depreciation & Amortization",
               33: "Interest", 34: "Total Expenses", 35: "Earnings Before Tax", 37: "Taxes", 38: "Net Earnings"}
    for r, l in is_rows.items():
        ws[f"A{r}"] = l
    for i, (c, h) in enumerate(zip(cols, list(hist) + [None] * n_f)):
        p = cols[i - 1] if i else None
        if h is not None:
            put(26, i, h.revenue, is_input=True); put(27, i, h.cogs, is_input=True)
            put(30, i, h.salaries, is_input=True); put(31, i, h.rent, is_input=True)
            put(32, i, h.da, is_input=True); put(33, i, h.interest, is_input=True); put(37, i, h.taxes, is_input=True)
        else:
            put(26, i, f"={p}26*(1+{c}8)"); put(27, i, f"={c}26*{c}9"); put(30, i, f"={c}10*{c}26")
            put(31, i, f"={c}11"); put(32, i, f"={c}67"); put(33, i, f"={c}105"); put(37, i, f"={c}35*{c}14")
        put(28, i, f"={c}26-{c}27"); put(34, i, f"=SUM({c}30:{c}33)"); put(35, i, f"={c}28-{c}34"); put(38, i, f"={c}35-{c}37")

    # balance sheet rows 44-60
    _section(ws, 41, "Balance Sheet")
    bs_rows = {43: "Assets", 44: "Cash", 45: "Accounts Receivable", 46: "Inventory", 47: "Property & Equipment",
               48: "Total Assets", 50: "Liabilities", 51: "Accounts Payable", 52: "Debt", 53: "Total Liabilities",
               54: "Shareholder's Equity", 55: "Equity Capital", 56: "Retained Earnings", 57: "Shareholder's Equity",
               58: "Total Liabilities & Shareholder's Equity", 60: "Check"}
    for r, l in bs_rows.items():
        ws[f"A{r}"] = l
    for i, (c, h) in enumerate(zip(cols, list(hist) + [None] * n_f)):
        p = cols[i - 1] if i else None
        if h is not None:
            put(44, i, h.cash, is_input=True); put(45, i, h.ar, is_input=True); put(46, i, h.inventory, is_input=True)
            put(47, i, h.ppe, is_input=True); put(51, i, h.ap, is_input=True); put(52, i, h.debt, is_input=True)
            put(55, i, h.equity_capital, is_input=True); put(56, i, h.retained_earnings, is_input=True)
        else:
            put(44, i, f"={c}82"); put(45, i, f"={c}26*{c}15/{c}21"); put(46, i, f"={c}27*{c}16/{c}21")
            put(47, i, f"={c}99"); put(51, i, f"={c}27*{c}17/{c}21"); put(52, i, f"={c}104")
            put(55, i, f"={p}55+{c}77"); put(56, i, f"={p}56+{c}38")
        put(48, i, f"=SUM({c}44:{c}47)"); put(53, i, f"=SUM({c}51:{c}52)"); put(57, i, f"=SUM({c}55:{c}56)")
        put(58, i, f"={c}53+{c}57"); put(60, i, f"={c}58-{c}48")

    # cash flow rows 66-82
    _section(ws, 63, "Cash Flow Statement")
    cf_rows = {65: "Operating Cash Flow", 66: "Net Earnings", 67: "Plus: Depreciation & Amortization",
               68: "Less: Changes in Working Capital", 69: "Cash from Operations", 71: "Investing Cash Flow",
               72: "Investments in Property & Equipment", 73: "Cash from Investing", 75: "Financing Cash Flow",
               76: "Issuance (repayment) of debt", 77: "Issuance (repayment) of equity", 78: "Cash from Financing",
               80: "Net Increase (decrease) in Cash", 81: "Opening Cash Balance", 82: "Closing Cash Balance",
               83: "Cash Check (vs Balance Sheet)"}
    for r, l in cf_rows.items():
        ws[f"A{r}"] = l
    for i, (c, h) in enumerate(zip(cols, list(hist) + [None] * n_f)):
        p = cols[i - 1] if i else None
        put(66, i, f"={c}38"); put(68, i, f"={c}93")
        if h is not None:
            put(67, i, f"={c}32"); put(72, i, h.capex, is_input=True)
            put(76, i, h.debt_issued, is_input=True); put(77, i, h.equity_issued, is_input=True)
        else:
            put(67, i, f"={c}98"); put(72, i, f"={c}97"); put(76, i, f"={c}103"); put(77, i, f"={c}20")
        put(69, i, f"={c}66+{c}67-{c}68"); put(73, i, f"=SUM({c}72)"); put(78, i, f"=SUM({c}76:{c}77)")
        put(80, i, f"={c}69-{c}73+{c}78"); put(81, i, 0 if i == 0 else f"={p}82", is_input=(i == 0))
        put(82, i, f"=SUM({c}80:{c}81)"); put(83, i, f"={c}82-{c}44")

    # schedules rows 89-105
    _section(ws, 86, "Supporting Schedules")
    sc_rows = {88: "Working Capital Schedule", 89: "Accounts Receivable", 90: "Inventory", 91: "Accounts Payable",
               92: "Net Working Capital (NWC)", 93: "Change in NWC", 95: "Depreciation Schedule", 96: "PPE Opening",
               97: "Plus Capex", 98: "Less Depreciation", 99: "PPE Closing", 101: "Debt & Interest Schedule",
               102: "Debt Opening", 103: "Issuance (repayment)", 104: "Debt Closing", 105: "Interest Expense"}
    for r, l in sc_rows.items():
        ws[f"A{r}"] = l
    for i, (c, h) in enumerate(zip(cols, list(hist) + [None] * n_f)):
        p = cols[i - 1] if i else None
        put(89, i, f"={c}45"); put(90, i, f"={c}46"); put(91, i, f"={c}51"); put(92, i, f"={c}89+{c}90-{c}91")
        put(93, i, f"={c}92" if i == 0 else f"={c}92-{p}92")
        put(96, i, ppe_opening0 if i == 0 else f"={p}99", is_input=(i == 0))
        put(102, i, debt_opening0 if i == 0 else f"={p}104", is_input=(i == 0))
        if h is not None:
            put(97, i, f"={c}72"); put(98, i, f"={c}32"); put(103, i, f"={c}76"); put(105, i, f"={c}33")
        else:
            put(97, i, f"={c}18"); put(98, i, f"={c}96*{c}12"); put(103, i, f"={c}19")
            put(105, i, f"=AVERAGE({c}102,{c}104)*{c}13")
        put(99, i, f"={c}96+{c}97-{c}98"); put(104, i, f"={c}102+{c}103")
    for c in cols:
        ws.column_dimensions[c].width = 13
    ws.freeze_panes = "B4"
    _legend(wb)
    path = Path(path)
    wb.save(path)
    return path


# --------------------------------------------------------------------------- DCF
def write_dcf(path: str | Path, inp: DCFInputs) -> Path:
    """Formula-driven DCF at the same coordinates as the CFI template (assumptions D6:D16, model rows 18-39)."""
    from .fin import to_date
    n = len(inp.ebit)
    wb = Workbook()
    ws = wb.active
    ws.title = "DCF Model"
    ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].width = 34
    _title(ws, "B3", "DCF Model")
    _section(ws, 5, "Assumptions", 2, 4)
    asm = [("Tax Rate", inp.tax_rate, PCT), ("Discount Rate", inp.discount_rate, PCT),
           ("Perpetual Growth Rate", inp.perpetual_growth, PCT), ("EV/EBITDA Multiple", inp.ev_ebitda_multiple, "0.0x"),
           ("Transaction Date", to_date(inp.transaction_date), "yyyy-mm-dd"), ("Fiscal Year End", to_date(inp.fiscal_year_end), "yyyy-mm-dd"),
           ("Current Price", inp.current_price, NUM2), ("Shares Outstanding", inp.shares_outstanding, NUM),
           ("Debt", inp.debt, NUM), ("Cash", inp.cash, NUM)]
    for k, (label, v, fmt) in enumerate(asm):
        ws[f"B{6+k}"] = label; c = ws[f"D{6+k}"]; c.value = v; c.number_format = fmt; c.font = BLUE
    capex = inp.capex_list()
    ws["B16"] = "Capex (per year, see row 25)"
    first = 5  # column E
    cols = [get_column_letter(first + t) for t in range(n)]
    exit_col = get_column_letter(first + n)          # J in the CFI template for n=5
    tv_lbl = get_column_letter(first + n + 2); tv_val = get_column_letter(first + n + 4)
    ws["B18"] = "Discounted Cash Flow"; ws["B18"].font = BOLD; ws["D18"] = "Entry"; ws[f"{exit_col}18"] = "Exit"
    ws["B19"] = "Date"; ws["B20"] = "Time Periods"; ws["B21"] = "Year Fraction"; ws["B22"] = "EBIT"
    ws["B23"] = "Less: Cash Taxes"; ws["B24"] = "Plus: D&A"; ws["B25"] = "Less: Capex"; ws["B26"] = "Less: Changes in NWC"
    ws["B27"] = "Unlevered FCF"; ws["B28"] = "(Entry)/Exit"; ws["B29"] = "Transaction CF (valuation)"; ws["B30"] = "Transaction CF (IRR)"
    ws["D19"] = "=D10"; ws["D19"].number_format = "yyyy-mm-dd"
    for t, c in enumerate(cols):
        p = cols[t - 1] if t else "D"
        ws[f"{c}18"] = f"=YEAR({c}19)"
        ws[f"{c}19"] = f"=DATE(YEAR($D$11)+{c}20,MONTH($D$11),DAY($D$11))"; ws[f"{c}19"].number_format = "yyyy-mm-dd"
        ws[f"{c}20"] = 0 if t == 0 else f"={p}20+1"
        ws[f"{c}21"] = f"=YEARFRAC({p}19,{c}19)"; ws[f"{c}21"].number_format = "0.00"
        for r, v in ((22, inp.ebit[t]), (24, inp.da[t]), (25, capex[t]), (26, inp.change_nwc[t])):
            ws[f"{c}{r}"] = v; ws[f"{c}{r}"].font = BLUE; ws[f"{c}{r}"].number_format = NUM
        ws[f"{c}23"] = f"={c}22*$D$6"; ws[f"{c}27"] = f"={c}22-{c}23+{c}24-{c}25-{c}26"
        ws[f"{c}29"] = f"=({c}28+{c}27)*{c}21"; ws[f"{c}30"] = f"=({c}28+{c}27)*{c}21"
        for r in (23, 27, 29, 30):
            ws[f"{c}{r}"].number_format = NUM
    last = cols[-1]
    ws[f"{exit_col}19"] = f"={last}19"; ws[f"{exit_col}19"].number_format = "yyyy-mm-dd"
    ws[f"{tv_lbl}18"] = "Terminal Value"; ws[f"{tv_lbl}18"].font = BOLD
    ws[f"{tv_lbl}19"] = "Perpetual Growth"; ws[f"{tv_val}19"] = f"=({last}27*(1+D8))/(D7-D8)"
    ws[f"{tv_lbl}20"] = "EV/EBITDA"; ws[f"{tv_val}20"] = f"=D9*({last}22+{last}24)"
    ws[f"{tv_lbl}21"] = "Average"; ws[f"{tv_val}21"] = f"=AVERAGE({tv_val}19:{tv_val}20)"
    used = {"average": f"={tv_val}21", "perpetuity": f"={tv_val}19", "multiple": f"={tv_val}20"}[inp.terminal_method]
    ws[f"{tv_lbl}22"] = f"Used ({inp.terminal_method})"; ws[f"{tv_val}22"] = used
    for r in (19, 20, 21, 22):
        ws[f"{tv_val}{r}"].number_format = NUM
    ws["D28"] = "=-I36".replace("I36", f"{tv_lbl}36"); ws[f"{exit_col}28"] = f"={tv_val}22"
    ws["D29"] = 0; ws[f"{exit_col}29"] = f"={exit_col}28+{exit_col}27"
    ws["D30"] = "=D28+D27"; ws[f"{exit_col}30"] = f"={exit_col}28"
    for cell in ("D28", "D29", "D30", f"{exit_col}28", f"{exit_col}29", f"{exit_col}30"):
        ws[cell].number_format = NUM
    # valuation blocks
    ws["B32"] = "Intrinsic Value"; ws["B32"].font = BOLD
    ws["B33"] = "Enterprise Value"; ws["D33"] = f"=XNPV(D7,D29:{exit_col}29,D19:{exit_col}19)"
    ws["B34"] = "Plus: Cash"; ws["D34"] = "=D15"; ws["B35"] = "Less: Debt"; ws["D35"] = "=D14"
    ws["B36"] = "Equity Value"; ws["D36"] = "=D33+D34-D35"; ws["B38"] = "Equity Value/Share"; ws["D38"] = "=D36/D13"
    mk = tv_lbl  # market block in the terminal-value label column
    ws[f"{mk}32"] = "Market Value"; ws[f"{mk}32"].font = BOLD
    ws[f"{mk}33"] = "Market Cap"; ws[f"{tv_val}33"] = "=D13*D12"; ws[f"{mk}34"] = "Plus: Debt"; ws[f"{tv_val}34"] = "=D14"
    ws[f"{mk}35"] = "Less: Cash"; ws[f"{tv_val}35"] = "=D15"; ws[f"{mk}36"] = "Enterprise Value"; ws[f"{tv_val}36"] = f"={tv_val}33+{tv_val}34-{tv_val}35"
    ws[f"{mk}38"] = "Equity Value/Share"; ws[f"{tv_val}38"] = "=D12"
    ws["D28"] = f"=-{tv_val}36"
    ws["B41"] = "Rate of Return"; ws["B41"].font = BOLD
    ws["B42"] = "Target Price Upside"; ws["D42"] = f"=D38/{tv_val}38-1"; ws["D42"].number_format = PCT
    ws["B43"] = "Internal Rate of Return (IRR)"; ws["D43"] = f"=XIRR(D30:{exit_col}30,D19:{exit_col}19)"; ws["D43"].number_format = PCT
    for cell in ("D33", "D34", "D35", "D36", f"{tv_val}33", f"{tv_val}34", f"{tv_val}35", f"{tv_val}36"):
        ws[cell].number_format = NUM
    for cell in ("D38", f"{tv_val}38"):
        ws[cell].number_format = NUM2
    for c in cols + [exit_col, tv_val]:
        ws.column_dimensions[c].width = 13
    ws.column_dimensions[tv_lbl].width = 26
    _legend(wb)
    path = Path(path)
    wb.save(path)
    return path


# --------------------------------------------------------------------------- generic value writers
def write_rows(ws, years: Sequence[Any], rows: Dict[str, Sequence[float]], start_row: int = 1, title: str | None = None,
               fmt: str = NUM) -> int:
    r = start_row
    if title:
        ws.cell(row=r, column=1, value=title).font = BOLD; r += 1
    for j, y in enumerate(years):
        c = ws.cell(row=r, column=2 + j, value=y); c.font = HEAD; c.fill = HEAD_FILL
    r += 1
    for label, vals in rows.items():
        if not isinstance(vals, (list, tuple)):
            continue
        ws.cell(row=r, column=1, value=label)
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=2 + j, value=v)
            c.number_format = PCT if (isinstance(v, float) and abs(v) < 1.5 and "ays" not in label and "eadcount" not in label and "Check" not in label) else fmt
        r += 1
    return r + 1


def write_projection(path: str | Path, res: Dict[str, Any]) -> Path:
    wb = Workbook(); ws = wb.active; ws.title = "Income Statement"
    years = res["years"]
    write_rows(ws, years, res["income_statement"], title="Income Statement (USD $000)")
    ws.column_dimensions["A"].width = 40
    ws2 = wb.create_sheet("Balance Sheet"); write_rows(ws2, years, res["balance_sheet"], title="Balance Sheet"); ws2.column_dimensions["A"].width = 40
    ws3 = wb.create_sheet("Cash Flow"); write_rows(ws3, years, res["cash_flow"], title="Cash Flow Statement"); ws3.column_dimensions["A"].width = 40
    ws4 = wb.create_sheet("Sales"); r = 1
    for name, s in res["sales"].items():
        r = write_rows(ws4, years, s, start_row=r, title=name)
    ws4.column_dimensions["A"].width = 30
    ws5 = wb.create_sheet("Payroll"); r = 1
    for t, s in res["payroll"]["types"].items():
        r = write_rows(ws5, years, {k: v for k, v in s.items() if isinstance(v, list)}, start_row=r, title=t)
    write_rows(ws5, years, {k: v for k, v in res["payroll"].items() if isinstance(v, list)}, start_row=r, title="Payroll totals")
    ws5.column_dimensions["A"].width = 30
    ws6 = wb.create_sheet("Opex"); write_rows(ws6, years, res["opex"], title="Operating expenses"); ws6.column_dimensions["A"].width = 36
    ws7 = wb.create_sheet("Schedules"); r = 1
    for name, s in res["schedules"].items():
        r = write_rows(ws7, years, {k: v for k, v in s.items() if isinstance(v, list)}, start_row=r, title=name)
    ws7.column_dimensions["A"].width = 30
    ws8 = wb.create_sheet("Ratios")
    from .ratios import flatten
    flat = [flatten(x) for x in res["ratios"]]
    rows = {k: [f[k] for f in flat] for k in flat[0]} if flat else {}
    write_rows(ws8, years, rows, title="Financial Ratio Analysis", fmt=NUM2); ws8.column_dimensions["A"].width = 44
    _legend(wb)
    path = Path(path); wb.save(path); return path


def write_generic(path: str | Path, sheets: Dict[str, Any]) -> Path:
    """sheets: name -> (years, rows-dict) written as values."""
    wb = Workbook(); first = True
    for name, (years, rows) in sheets.items():
        ws = wb.active if first else wb.create_sheet(); first = False
        ws.title = name[:31]
        r = 1
        if rows and all(isinstance(v, dict) for v in rows.values()):
            for sub, subrows in rows.items():
                r = write_rows(ws, years, {k: v for k, v in subrows.items() if isinstance(v, list)}, start_row=r, title=sub)
        else:
            write_rows(ws, years, {k: v for k, v in rows.items() if isinstance(v, list)}, title=name)
        ws.column_dimensions["A"].width = 40
    _legend(wb); path = Path(path); wb.save(path); return path


def _legend(wb):
    ws = wb.create_sheet("README")
    ws["A1"] = "Generated by financial-analysis-toolkit (finmodel)"; ws["A1"].font = BOLD
    ws["A3"] = "Blue = hard-coded input"; ws["A3"].font = BLUE
    ws["A4"] = "Black = formula"
    ws["A6"] = "Layout follows the CFI (Corporate Finance Institute) template conventions; engines re-implemented in Python."
    ws.column_dimensions["A"].width = 90


# --------------------------------------------------------------------------- LibreOffice recalc
def soffice_path() -> Optional[str]:
    for cand in ("soffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice", "/opt/homebrew/bin/soffice"):
        p = shutil.which(cand) or (cand if Path(cand).exists() else None)
        if p:
            return p
    return None


def recalc_with_libreoffice(path: str | Path) -> Optional[Path]:
    """Round-trip through LibreOffice headless so every formula gets a cached value; returns new file path."""
    exe = soffice_path()
    if not exe:
        return None
    path = Path(path)
    out_dir = Path(tempfile.mkdtemp(prefix="finmodel_recalc_"))
    subprocess.run([exe, "--headless", "--convert-to", "xlsx", "--outdir", str(out_dir), str(path)],
                   check=False, capture_output=True, timeout=240)
    out = out_dir / path.name
    return out if out.exists() else None


def read_values(path: str | Path, sheet: str) -> Dict[str, Any]:
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
