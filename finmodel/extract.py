"""Extract the structure of any Excel financial template into a JSON + Markdown spec.

For every sheet we record each non-empty cell (cached value + formula), classify cells as
input (hard-coded number), formula or label, relativise formulas to R1C1-style offsets so
repeating row patterns collapse to one signature, and list cross-sheet links.  This is what
was used to reverse-engineer the CFI templates into the engines in this package.

.xls files are converted to .xlsx with LibreOffice (soffice) when available, otherwise
read values-only with xlrd.
"""
from __future__ import annotations

import datetime as _dt
import json
import re
import shutil
import subprocess
import tempfile
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

# A1 reference not preceded by sheet-quote/identifier chars and not followed by '(' (function names)
_REF = re.compile(r"(?<![A-Za-z0-9_.])(\$?)([A-Z]{1,3})(\$?)(\d{1,7})(?![A-Za-z0-9_(])")
_SHEET = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_.]+))!")


def relativize(formula: str, row: int, col: int) -> str:
    """'=B5*C$3' at C5 -> '=R[0]C[-1]*R3C[0]'  (absolute parts keep their number)."""
    def rep(m):
        cabs, cl, rabs, rw = m.group(1), m.group(2), m.group(3), int(m.group(4))
        c = column_index_from_string(cl)
        rpart = f"R{rw}" if rabs else f"R[{rw - row}]"
        cpart = f"C{c}" if cabs else f"C[{c - col}]"
        return rpart + cpart
    return _REF.sub(rep, formula)


def _soffice() -> Optional[str]:
    for cand in ("soffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice", "/opt/homebrew/bin/soffice"):
        p = shutil.which(cand) or (cand if Path(cand).exists() else None)
        if p:
            return p
    return None


def convert_xls_to_xlsx(path: Path, out_dir: Path) -> Optional[Path]:
    exe = _soffice()
    if not exe:
        return None
    subprocess.run([exe, "--headless", "--convert-to", "xlsx", "--outdir", str(out_dir), str(path)],
                   check=False, capture_output=True, timeout=180)
    out = out_dir / (path.stem + ".xlsx")
    return out if out.exists() else None


def _jsonable(v):
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.isoformat()
    return v


def _extract_xls_values_only(path: Path) -> Dict[str, Any]:
    import xlrd  # optional dependency
    book = xlrd.open_workbook(str(path))
    sheets = []
    for sh in book.sheets():
        cells = {}
        for r in range(sh.nrows):
            for c in range(sh.ncols):
                v = sh.cell_value(r, c)
                if v not in ("", None):
                    cells[f"{get_column_letter(c + 1)}{r + 1}"] = {"v": v, "f": None, "kind": "input" if isinstance(v, (int, float)) else "label"}
        sheets.append({"name": sh.name, "dims": f"A1:{get_column_letter(max(sh.ncols, 1))}{max(sh.nrows, 1)}",
                       "cells": cells, "rows": [], "stats": {"cells": len(cells), "formulas": 0, "inputs": sum(1 for x in cells.values() if x["kind"] == "input")},
                       "cross_sheet_refs": [], "formula_patterns": []})
    return {"file": str(path), "format": "xls (values only)", "sheets": sheets, "defined_names": []}


def extract_workbook(path: str | Path, out_dir: str | Path | None = None, max_cells_per_sheet: int = 20000) -> Dict[str, Any]:
    path = Path(path)
    tmp = None
    src = path
    if path.suffix.lower() == ".xls":
        tmp = Path(tempfile.mkdtemp(prefix="finmodel_xls_"))
        conv = convert_xls_to_xlsx(path, tmp)
        if conv is None:
            spec = _extract_xls_values_only(path)
            if out_dir:
                _write(spec, path, Path(out_dir))
            return spec
        src = conv
    wb_f = openpyxl.load_workbook(src, data_only=False)
    wb_v = openpyxl.load_workbook(src, data_only=True)
    sheets: List[Dict[str, Any]] = []
    for ws in wb_f.worksheets:
        wv = wb_v[ws.title]
        cells: Dict[str, Any] = {}
        rows: Dict[int, Dict[str, Any]] = {}
        xrefs: Counter = Counter()
        patterns: Counter = Counter()
        n = 0
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                n += 1
                if n > max_cells_per_sheet:
                    break
                v = c.value
                is_formula = isinstance(v, str) and v.startswith("=")
                cached = _jsonable(wv[c.coordinate].value)
                if is_formula:
                    kind = "formula"
                    pat = relativize(v, c.row, c.column)
                    patterns[pat] += 1
                    for m in _SHEET.finditer(v):
                        name = (m.group(1) or m.group(2)).replace("''", "'")
                        if name != ws.title:
                            xrefs[name] += 1
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    kind = "input"
                elif isinstance(v, (_dt.datetime, _dt.date)):
                    kind = "date"
                else:
                    kind = "label"
                cells[c.coordinate] = {"v": cached, "f": v if is_formula else None, "kind": kind}
                r = rows.setdefault(c.row, {"row": c.row, "label": None, "inputs": [], "formula": None, "pattern": None})
                if kind == "label" and r["label"] is None and c.column <= 3:
                    r["label"] = str(v).strip()
                elif kind == "input":
                    r["inputs"].append({"cell": c.coordinate, "v": v})
                elif kind == "formula" and r["formula"] is None:
                    r["formula"] = v
                    r["pattern"] = relativize(v, c.row, c.column)
        stats = {"cells": len(cells),
                 "formulas": sum(1 for x in cells.values() if x["kind"] == "formula"),
                 "inputs": sum(1 for x in cells.values() if x["kind"] == "input"),
                 "labels": sum(1 for x in cells.values() if x["kind"] == "label"),
                 "unique_formula_patterns": len(patterns)}
        sheets.append({"name": ws.title, "dims": ws.dimensions, "stats": stats,
                       "cross_sheet_refs": [{"sheet": k, "refs": v} for k, v in xrefs.most_common()],
                       "formula_patterns": [{"pattern": k, "count": v} for k, v in patterns.most_common(60)],
                       "rows": [rows[k] for k in sorted(rows)], "cells": cells})
    names = []
    try:
        for name, dn in wb_f.defined_names.items():
            names.append({"name": name, "refers_to": dn.attr_text})
    except AttributeError:  # older openpyxl
        for dn in wb_f.defined_names.definedName:
            names.append({"name": dn.name, "refers_to": dn.attr_text})
    spec = {"file": str(path), "format": "xlsx", "sheets": sheets, "defined_names": names,
            "totals": {k: sum(s["stats"][k] for s in sheets) for k in ("cells", "formulas", "inputs")}}
    if tmp:
        shutil.rmtree(tmp, ignore_errors=True)
    if out_dir:
        _write(spec, path, Path(out_dir))
    return spec


def to_markdown(spec: Dict[str, Any], max_rows: int = 400) -> str:
    out = [f"# {Path(spec['file']).name}", "", f"Format: {spec['format']}  ",
           f"Totals: {spec.get('totals', {})}", ""]
    for s in spec["sheets"]:
        out += [f"## Sheet: {s['name']}  (`{s['dims']}`)", "", f"Stats: {s['stats']}", ""]
        if s["cross_sheet_refs"]:
            out += ["Links to: " + ", ".join(f"`{x['sheet']}` ({x['refs']})" for x in s["cross_sheet_refs"]), ""]
        if s["rows"]:
            out += ["| row | label | first formula | pattern | inputs |", "|---|---|---|---|---|"]
            for r in s["rows"][:max_rows]:
                inputs = ", ".join(f"{i['cell']}={i['v']}" for i in r["inputs"][:6])
                if len(r["inputs"]) > 6:
                    inputs += f" (+{len(r['inputs']) - 6})"
                f = (r["formula"] or "").replace("|", "\\|")
                p = (r["pattern"] or "").replace("|", "\\|")
                out.append(f"| {r['row']} | {r['label'] or ''} | `{f}` | `{p}` | {inputs} |")
            out.append("")
    return "\n".join(out)


def _write(spec: Dict[str, Any], path: Path, out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = path.stem
    (out_dir / f"{stem}.spec.json").write_text(json.dumps(spec, indent=1, default=str))
    (out_dir / f"{stem}.spec.md").write_text(to_markdown(spec))
