import pytest
from pathlib import Path
from finmodel import extract
from conftest import cfi_file


def test_relativize():
    assert extract.relativize("=B5*C$3+$A$1", 5, 3) == "=R[0]C[-1]*R3C[0]+R1C1"
    assert extract.relativize("=SUM(D30:D33)", 34, 4) == "=SUM(R[-4]C[0]:R[-1]C[0])"
    assert extract.relativize("='Payroll 2018'!N15", 10, 2) == "='Payroll 2018'!R[5]C[12]"
    assert extract.relativize("=YEAR(E19)", 18, 5) == "=YEAR(R[1]C[0])"
    assert "DATE(" in extract.relativize("=DATE(YEAR($D$11)+E20,6,30)", 19, 5)


def test_extract_synthetic(tmp_path):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Model"
    ws["A1"] = "Revenue"; ws["B1"] = 100; ws["C1"] = "=B1*1.1"
    ws2 = wb.create_sheet("Out"); ws2["A1"] = "Total"; ws2["B1"] = "=Model!C1"
    p = tmp_path / "t.xlsx"; wb.save(p)
    spec = extract.extract_workbook(p, tmp_path / "spec")
    s = {x["name"]: x for x in spec["sheets"]}
    assert s["Model"]["stats"] == {"cells": 3, "formulas": 1, "inputs": 1, "labels": 1, "unique_formula_patterns": 1}
    assert s["Out"]["cross_sheet_refs"] == [{"sheet": "Model", "refs": 1}]
    assert s["Model"]["rows"][0]["label"] == "Revenue" and s["Model"]["rows"][0]["pattern"] == "=R[0]C[-1]*1.1"
    assert (tmp_path / "spec" / "t.spec.json").exists() and (tmp_path / "spec" / "t.spec.md").exists()


@pytest.mark.skipif(cfi_file("CFI-DCF-Model-Template-Updated.xlsx") is None, reason="CFI template not on this machine")
def test_extract_cfi_dcf():
    spec = extract.extract_workbook(cfi_file("CFI-DCF-Model-Template-Updated.xlsx"))
    dcf = next(s for s in spec["sheets"] if s["name"] == "DCF Model")
    assert dcf["stats"]["formulas"] >= 60
    labels = {r["label"] for r in dcf["rows"] if r["label"]}
    assert {"Unlevered FCF", "Enterprise Value", "Equity Value/Share"} <= labels
