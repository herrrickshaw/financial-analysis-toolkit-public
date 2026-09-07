"""LBO and merger engines: reconciliation to the source workbooks (when downloaded) + identities."""
import json
from pathlib import Path
import pytest
from finmodel import lbo, merger

ROOT = Path(__file__).resolve().parent.parent
ASM = ROOT / "downloads/asimplemodel/LBO_Simple LBO Scenarios and Data Tables.xlsx"
BIWS = ROOT / "downloads/biws/108-11-Merger-Model-Interview-Questions.xlsx"


def test_lbo_identities_and_example():
    d = json.loads((ROOT / "examples/lbo_asm.json").read_text())
    r = lbo.from_dict(d)
    assert r["balanced"]
    assert abs(r["closing_balance_sheet"]["check"]) < 1e-6
    s, u = r["sources"], r["uses"]
    assert s["total"] == pytest.approx(u["total"]) == pytest.approx(100000 + 1000 + 1000)
    assert s["Subordinated Debt"] == pytest.approx(2.5 * (108115.46064 - 81086.59548 - 11892.7006704))
    IS, CF, BS = r["income_statement"], r["cash_flow"], r["balance_sheet"]
    assert BS["Cash"] == pytest.approx(CF["Ending Cash"])
    assert all(abs(a - b) < 1e-9 for a, b in zip(BS["Senior Debt"], r["debt_schedule"]["Senior Debt"]["Ending"]))
    assert r["debt_schedule"]["Senior Debt"]["Ending"][-1] == pytest.approx(0.0)      # 5-year straight-line
    assert r["returns"]["sponsor"]["irr"] == pytest.approx(0.30687164664268496, abs=1e-7)
    assert r["returns"]["Subordinated Debt"]["irr"] == pytest.approx(0.17153204083442686, abs=1e-7)
    assert r["sensitivity"]["table"]["Scenario 3"][0] == pytest.approx(0.05220653116703035, abs=1e-7)


def test_lbo_sweep_and_pik_mechanics():
    d = json.loads((ROOT / "examples/lbo_asm.json").read_text())
    d["tranches"][1]["sweep_pct"] = 1.0          # senior debt prepaid with all excess cash
    d["tranches"][0]["pik_pct"] = 0.5            # half of sub-debt interest accrues
    d.pop("scenarios"); d.pop("exit_multiples")
    r = lbo.from_dict(d)
    assert r["balanced"]
    sen = r["debt_schedule"]["Senior Debt"]; sub = r["debt_schedule"]["Subordinated Debt"]
    assert sum(sen["Cash sweep"]) > 0 and sen["Ending"][-1] == pytest.approx(0.0)
    assert all(x >= -1e-9 for x in r["balance_sheet"]["Cash"])                # sweep never breaches minimum cash
    assert sub["Ending"][-1] > sub["Beginning"][0]                            # PIK accretes the balance
    assert r["cash_flow"]["PIK interest"][0] == pytest.approx(sub["PIK accrual"][0])


@pytest.mark.skipif(not ASM.exists(), reason="A Simple Model workbook not downloaded")
def test_lbo_reconciles_to_asm_workbook():
    import openpyxl
    wv = openpyxl.load_workbook(ASM, data_only=True)["LBO"]; v = lambda a: wv[a].value
    d = json.loads((ROOT / "examples/lbo_asm.json").read_text()); r = lbo.from_dict(d)
    for i, c in enumerate("FGHIJ"):
        assert r["income_statement"]["Net Income"][i] == pytest.approx(v(f"{c}53"), abs=1e-6)
        assert r["balance_sheet"]["Total Assets"][i] == pytest.approx(v(f"{c}74"), abs=1e-6)
        assert r["balance_sheet"]["Cash"][i] == pytest.approx(v(f"{c}60"), abs=1e-6)
    assert r["returns"]["exit_equity_value"] == pytest.approx(v("J205"), abs=1e-6)
    for si, row in enumerate([246, 247, 248]):
        for ci, col in enumerate("DEFGHIJ"):
            assert r["sensitivity"]["table"][f"Scenario {si+1}"][ci] == pytest.approx(v(f"{col}{row}"), abs=1e-7)


def test_merger_deal_reconciles_to_biws_numbers():
    d = json.loads((ROOT / "examples/merger_biws.json").read_text())
    r = merger.from_dict(d)["deal"]
    assert r["acquirer"]["eps"] == pytest.approx(1.0) and r["target"]["eps"] == pytest.approx(2.0)
    assert r["purchase_equity_value"] == pytest.approx(150) and r["purchase_enterprise_value"] == pytest.approx(150)
    assert r["new_shares_issued"] == pytest.approx(6) and r["combined_shares"] == pytest.approx(26)
    assert r["combined"]["net_income"] == pytest.approx(30) and r["combined"]["eps"] == pytest.approx(30 / 26)
    assert r["accretion_dilution_pct"] == pytest.approx(30 / 26 - 1)
    assert r["combined"]["ev_ebitda"] == pytest.approx((500 + 150) / 72)


@pytest.mark.skipif(not BIWS.exists(), reason="BIWS workbook not downloaded")
def test_merger_deal_reconciles_to_biws_workbook():
    import openpyxl
    ws = openpyxl.load_workbook(BIWS, data_only=True).active; v = lambda a: ws[a].value
    r = merger.from_dict(json.loads((ROOT / "examples/merger_biws.json").read_text()))["deal"]
    for key, cell in [("purchase_enterprise_value", "F38"), ("new_shares_issued", "F43"), ("accretion_dilution_pct", "F59")]:
        assert r[key] == pytest.approx(v(cell), abs=1e-9)
    assert r["combined"]["pe"] == pytest.approx(v("F55"), abs=1e-9)


def test_merger_pro_forma_identities():
    d = json.loads((ROOT / "examples/merger_biws.json").read_text())
    pf = merger.from_dict(d)["pro_forma"]
    ppa = pf["purchase_price_allocation"]
    assert ppa["goodwill"] == pytest.approx(900 - 300 + 50 - 100 - 200 + 300 * 0.25)
    assert pf["new_shares_issued"] == pytest.approx(900 * 0.3 / 30)
    ds = pf["debt_schedule"]
    assert ds["Opening"][0] == pytest.approx(900 * 0.4) and ds["Closing"][-1] == pytest.approx(0.0)   # balloon at maturity
    IS = pf["income_statement"]
    assert IS["Amortization of new intangibles"][0] == pytest.approx(-40) and IS["Depreciation of PP&E write-up"][0] == pytest.approx(-10)
    assert IS["Cost synergies (realised)"] == pytest.approx([20, 40, 40])
    assert IS["Transaction fees"] == pytest.approx([-15, 0, 0])
    assert IS["Pro forma shares"][0] == pytest.approx(109)
    # adjusted EPS strips one-offs and write-up D&A
    assert IS["Adjusted EPS (ex one-offs & write-up D&A)"][2] > IS["Pro forma EPS"][2]
    sens = merger.from_dict(d)["deal_sensitivity"]
    assert len(sens["table"]) == 5 and sens["table"][0][4] > sens["table"][4][4]   # lower premium, more accretive
