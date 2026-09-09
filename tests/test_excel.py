"""Write formula workbooks and prove (via LibreOffice recalculation) that they reproduce the engines."""
import pytest
from finmodel import three_statement as ts, dcf, projection
from finmodel import excel


def test_write_three_statement_and_dcf_have_formulas(tmp_path, three_statement_inputs, dcf_inputs):
    d = three_statement_inputs
    hist = [ts.HistoricalYear(**h) for h in d["historical"]]
    p = excel.write_three_statement(tmp_path / "ts.xlsx", hist, ts.ForecastAssumptions(**d["forecast"]), d["ppe_opening0"], d["debt_opening0"])
    from openpyxl import load_workbook
    ws = load_workbook(p)["Three Statement Model"]
    assert ws["I26"].value == "=H26*(1+I8)" and ws["M60"].value == "=M58-M48"
    q = excel.write_dcf(tmp_path / "dcf.xlsx", dcf.DCFInputs(**{k: v for k, v in dcf_inputs.items() if not k.startswith("_")}))
    wd = load_workbook(q)["DCF Model"]
    assert wd["D33"].value.startswith("=XNPV(") and wd["D43"].value.startswith("=XIRR(")


@pytest.mark.skipif(excel.soffice_path() is None, reason="LibreOffice not installed")
def test_libreoffice_recalc_matches_engine(tmp_path, three_statement_inputs, dcf_inputs):
    d = three_statement_inputs
    hist = [ts.HistoricalYear(**h) for h in d["historical"]]
    fc = ts.ForecastAssumptions(**d["forecast"])
    res = ts.run(hist, fc, d["ppe_opening0"], d["debt_opening0"])
    p = excel.write_three_statement(tmp_path / "ts.xlsx", hist, fc, d["ppe_opening0"], d["debt_opening0"], res)
    rp = excel.recalc_with_libreoffice(p)
    assert rp is not None
    v = excel.read_values(rp, "Three Statement Model")
    cols = "DEFGHIJKLM"
    for i, c in enumerate(cols):
        assert v[f"{c}38"] == pytest.approx(res.rows["Net Earnings"][i], abs=0.01)
        assert v[f"{c}48"] == pytest.approx(res.rows["Total Assets"][i], abs=0.01)
        assert v[f"{c}82"] == pytest.approx(res.rows["Closing Cash Balance"][i], abs=0.01)
        assert abs(v.get(f"{c}60", 0) or 0) < 0.01
        assert v[f"{c}3"] == "OK"
    inp = dcf.DCFInputs(**{k: v_ for k, v_ in dcf_inputs.items() if not k.startswith("_")})
    r = dcf.run(inp)
    q = excel.write_dcf(tmp_path / "dcf.xlsx", inp)
    rq = excel.recalc_with_libreoffice(q)
    w = excel.read_values(rq, "DCF Model")
    assert w["D33"] == pytest.approx(r["enterprise_value"], rel=1e-6)
    assert w["D38"] == pytest.approx(r["equity_value_per_share"], rel=1e-6)
    assert w["D43"] == pytest.approx(r["irr"], abs=1e-5)
    assert w["E21"] == pytest.approx(0.5)


def test_write_projection(tmp_path, projection_inputs):
    r = projection.from_dict(projection_inputs)
    p = excel.write_projection(tmp_path / "proj.xlsx", r)
    from openpyxl import load_workbook
    wb = load_workbook(p)
    assert {"Income Statement", "Balance Sheet", "Cash Flow", "Sales", "Payroll", "Opex", "Schedules", "Ratios"} <= set(wb.sheetnames)


def test_write_record_tables_writes_headers_rows_and_readme(tmp_path):
    sheets = {
        "States": {"headers": ["State", "Rate"], "rows": [["Gujarat", 0.75], ["Odisha", 1.0]]},
        "Central": {"headers": ["Scheme", "Outlay"], "rows": [["PLI Auto", 25938]]},
    }
    p = excel.write_record_tables(tmp_path / "catalog.xlsx", sheets, readme_lines=["Generated for testing", "Line 2"])
    from openpyxl import load_workbook
    wb = load_workbook(p)
    assert wb.sheetnames == ["README", "States", "Central"]
    assert wb["README"]["A1"].value == "Generated for testing"
    assert wb["README"]["A2"].value == "Line 2"
    ws = wb["States"]
    assert [ws.cell(row=1, column=c).value for c in (1, 2)] == ["State", "Rate"]
    assert [ws.cell(row=2, column=c).value for c in (1, 2)] == ["Gujarat", 0.75]
    assert [ws.cell(row=3, column=c).value for c in (1, 2)] == ["Odisha", 1.0]
    assert ws.freeze_panes == "A2"
    ws2 = wb["Central"]
    assert ws2.cell(row=2, column=1).value == "PLI Auto"
    assert ws2.cell(row=2, column=2).value == 25938


def test_write_record_tables_without_readme_uses_first_sheet_as_active(tmp_path):
    sheets = {"Only": {"headers": ["A"], "rows": [[1], [2]]}}
    p = excel.write_record_tables(tmp_path / "no_readme.xlsx", sheets)
    from openpyxl import load_workbook
    wb = load_workbook(p)
    assert wb.sheetnames == ["Only"]
    assert wb["Only"].cell(row=3, column=1).value == 2
